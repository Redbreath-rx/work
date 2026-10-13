import csv
import io
import json
import os
import re
import secrets
import uuid
import zipfile
from datetime import datetime, timezone
from functools import wraps
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit

from argon2 import PasswordHasher
from argon2.exceptions import VerificationError, InvalidHashError
from bson import ObjectId
from bson.errors import InvalidId
from flask import (Flask, abort, flash, g, redirect, render_template, request,
                   send_from_directory, session, url_for)
from flask_admin import Admin, AdminIndexView, expose
from flask_admin.contrib.pymongo import ModelView
from flask_wtf import CSRFProtect, FlaskForm
from openpyxl import load_workbook
from PIL import Image
from pymongo import MongoClient
from pymongo.errors import DuplicateKeyError, PyMongoError
from wtforms import PasswordField, SelectField, StringField, TextAreaField
from wtforms.validators import DataRequired, Length
from database import TYPES, initialize
from processing import input_hash, process, validate_output

ROOT = Path(__file__).resolve().parent
HASHER = PasswordHasher()


def now():
    return datetime.now(timezone.utc)


def split_values(value):
    if isinstance(value, list):
        value = ','.join(str(v) for v in value)
    return list(dict.fromkeys(x.strip() for x in re.split(r'[,，;；\n]', value or '') if x.strip()))


def clean_content(value):
    value = (value or '').strip()
    if not re.search(r'</?[A-Za-z][^>]*>', value):
        return value

    class TextParser(HTMLParser):
        def __init__(self):
            super().__init__(convert_charrefs=True)
            self.parts = []
            self.skip = 0

        def handle_starttag(self, tag, attrs):
            if tag in ('script', 'style'):
                self.skip += 1
            if not self.skip and tag in ('p', 'div', 'br', 'li'):
                self.parts.append('\n')

        def handle_endtag(self, tag):
            if tag in ('script', 'style') and self.skip:
                self.skip -= 1
            elif not self.skip and tag in ('p', 'div', 'li'):
                self.parts.append('\n')

        def handle_data(self, data):
            if not self.skip:
                self.parts.append(data)

    parser = TextParser()
    parser.feed(value)
    return re.sub(r'\n{3,}', '\n\n', ''.join(parser.parts)).strip()


def canonical_url(value):
    value = (value or '').strip()
    try:
        parsed = urlsplit(value)
        if parsed.scheme not in ('http', 'https') or not parsed.hostname or parsed.username:
            raise ValueError
        parsed.port  # Reject invalid ports.
    except ValueError:
        raise ValueError('来源网址须为有效的HTTP或HTTPS网址。') from None
    return urlunsplit((parsed.scheme.lower(), parsed.netloc.lower(),
                       parsed.path, parsed.query, ''))


def oid(value):
    try:
        return ObjectId(value)
    except (InvalidId, TypeError):
        abort(404)


def admin_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not g.user:
            return redirect(url_for('login'))
        if g.user['role'] != 'admin':
            abort(403)
        return fn(*args, **kwargs)
    return wrapper


def login_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not g.user:
            return redirect(url_for('login'))
        return fn(*args, **kwargs)
    return wrapper


class StoryForm(FlaskForm):
    title = StringField('故事名称', validators=[DataRequired(), Length(max=200)])
    content_kind = SelectField('内容类型', choices=['故事正文', '项目介绍'])
    story_type = SelectField('故事类型', choices=['待分类'] + TYPES)
    original_content = TextAreaField('原始内容', validators=[DataRequired()])
    summary = TextAreaField('摘要（最多200字）', validators=[Length(max=200)])
    regions = StringField('地域（逗号分隔）')
    ethnic_groups = StringField('相关民族（逗号分隔）')
    main_characters = StringField('主要人物（逗号分隔）')
    theme_keywords = StringField('主题关键词（逗号分隔）')
    cultural_meaning = TextAreaField('文化意义')
    site_name = StringField('来源网站', validators=[DataRequired()])
    source_url = StringField('来源网址', validators=[DataRequired()])


def story_from_form(data, previous=None):
    previous = previous or {}
    if data.get('content_kind', '故事正文') not in ['故事正文', '项目介绍']:
        raise ValueError('内容类型不正确。')
    kind = data.get('story_type', '待分类')
    if kind not in TYPES + ['待分类']:
        raise ValueError('故事类型不正确。')
    title = (data.get('title') or '').strip()
    content = clean_content(data.get('original_content'))
    site = (data.get('site_name') or '').strip()
    summary = (data.get('summary') or '').strip()
    if not title or not content or not site:
        raise ValueError('名称、原始内容和来源网站不能为空。')
    if len(title) > 200 or len(summary) > 200:
        raise ValueError('名称或摘要超过200字。')
    doc = {'title': title, 'content_kind': data.get('content_kind', '故事正文'),
           'story_type': kind, 'original_content': content, 'summary': summary,
           'source': {'site_name': site, 'url': canonical_url(data.get('source_url')),
                      'collected_at': previous.get('source', {}).get('collected_at', now())},
           'cultural_meaning': (data.get('cultural_meaning') or '').strip(),
           'images': previous.get('images', []),
           'status': previous.get('status', 'draft'),
           'created_at': previous.get('created_at', now()), 'updated_at': now()}
    for field in ['regions', 'ethnic_groups', 'main_characters', 'theme_keywords']:
        doc[field] = split_values(data.get(field))
    changed = any(previous.get(k) != doc[k] for k in
                  ['title', 'content_kind', 'original_content'])
    if changed or not summary or kind == '待分类':
        doc['status'] = 'draft'
    # Editing a deleted resource must not silently restore it.
    if previous.get('status') == 'deleted':
        doc['status'] = 'deleted'
    return doc


def create_app(config=None):
    app = Flask(__name__)
    runtime = ROOT / '.runtime'
    runtime.mkdir(exist_ok=True)
    keyfile = runtime / 'session-key'
    if not keyfile.exists():
        fd = os.open(keyfile, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
        with os.fdopen(fd, 'w') as handle:
            handle.write(secrets.token_hex(32))
    app.config.update(
        SECRET_KEY=os.environ.get('SESSION_SECRET') or keyfile.read_text(),
        MONGO_URI=os.environ.get('MONGO_URI', 'mongodb://127.0.0.1:27017/?replicaSet=rs0'),
        MONGO_DB=os.environ.get('MONGO_DB', 'folk_stories'),
        MAX_CONTENT_LENGTH=8 * 1024 * 1024,
        SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE='Lax',
        UPLOAD_FOLDER=str(ROOT / 'uploads'),
    )
    if config:
        app.config.update(config)
    client = MongoClient(app.config['MONGO_URI'], serverSelectionTimeoutMS=5000, tz_aware=True)
    db = client[app.config['MONGO_DB']]
    app.extensions['mongo_client'] = client
    app.extensions['mongo_db'] = db
    Path(app.config['UPLOAD_FOLDER']).mkdir(exist_ok=True)
    CSRFProtect(app)

    @app.before_request
    def load_user():
        g.db = db
        user_id = session.get('user_id')
        g.user = db.users.find_one({'_id': oid(user_id)}) if user_id else None

    @app.context_processor
    def context():
        return {'types': TYPES, 'current_user': getattr(g, 'user', None),
                'status_labels': {'draft': '草稿', 'published': '已发布', 'deleted': '已删除'},
                'review_labels': {'pending': '待审核', 'accepted': '已采用', 'rejected': '已拒绝'},
                'llm_enabled': bool(os.environ.get('AI_API_KEY'))}

    @app.errorhandler(403)
    @app.errorhandler(404)
    @app.errorhandler(400)
    @app.errorhandler(413)
    def error_page(error):
        return render_template('error.html', code=error.code,
                               message={403: '你没有权限执行此操作。', 404: '资源不存在或尚未发布。',
                                        400: '请求无效，请刷新页面重试。',
                                        413: '文件不能超过8 MB。'}.get(error.code, '请求失败。')), error.code

    @app.get('/health')
    def health():
        try:
            client.admin.command('ping')
            return {'status': 'ok', 'database': 'MongoDB',
                    'processing': 'api' if os.environ.get('AI_API_KEY') else 'local-rules'}
        except PyMongoError:
            return {'status': 'unavailable'}, 503

    @app.get('/')
    def index():
        q = request.args.get('q', '').strip()[:100]
        kind = request.args.get('type', '')
        region = request.args.get('region', '').strip()
        keyword = request.args.get('keyword', '').strip()
        query = {'status': 'published'}
        if q:
            query['title'] = {'$regex': re.escape(q), '$options': 'i'}
        if kind in TYPES:
            query['story_type'] = kind
        if region:
            query['regions'] = region
        if keyword:
            query['theme_keywords'] = keyword
        try:
            page = max(1, int(request.args.get('page', 1)))
        except ValueError:
            page = 1
        total = db.stories.count_documents(query)
        pages = max(1, (total + 8) // 9)
        page = min(page, pages)
        stories = list(db.stories.find(query).sort('_id', 1).skip((page - 1) * 9).limit(9))
        return render_template('index.html', stories=stories, total=total, page=page,
                               pages=pages, q=q, kind=kind, region=region, keyword=keyword)

    @app.get('/stories/<story_id>')
    def detail(story_id):
        story = db.stories.find_one({'_id': oid(story_id), 'status': 'published'})
        if not story:
            abort(404)
        favorite = bool(g.user and db.favorites.find_one(
            {'user_id': g.user['_id'], 'story_id': story['_id']}))
        return render_template('detail.html', story=story, favorite=favorite)

    @app.route('/login', methods=['GET', 'POST'])
    def login():
        if request.method == 'POST':
            user = db.users.find_one({'username': request.form.get('username', '').strip()})
            valid = False
            if user:
                try:
                    valid = HASHER.verify(user['password_hash'], request.form.get('password', ''))
                except (VerificationError, InvalidHashError):
                    pass
            if valid:
                session.clear()
                session['user_id'] = str(user['_id'])
                flash('登录成功。', 'success')
                return redirect(url_for('manage' if user['role'] == 'admin' else 'index'))
            flash('用户名或密码错误。', 'error')
        return render_template('auth.html', register=False)

    @app.route('/register', methods=['GET', 'POST'])
    def register():
        if request.method == 'POST':
            username = request.form.get('username', '').strip()
            password = request.form.get('password', '')
            if not re.fullmatch(r'[\w\u4e00-\u9fff-]{2,30}', username) or not 8 <= len(password) <= 128:
                flash('用户名须为2—30个字母、数字、汉字或下划线；密码须为8—128位。', 'error')
            else:
                try:
                    db.users.insert_one({'username': username, 'password_hash': HASHER.hash(password),
                                         'role': 'user', 'created_at': now()})
                    flash('注册成功，请登录。', 'success')
                    return redirect(url_for('login'))
                except DuplicateKeyError:
                    flash('用户名已存在。', 'error')
        return render_template('auth.html', register=True)

    @app.post('/logout')
    def logout():
        session.clear()
        return redirect(url_for('index'))

    @app.post('/stories/<story_id>/favorite')
    @login_required
    def toggle_favorite(story_id):
        identifier = oid(story_id)
        existing = db.favorites.find_one({'user_id': g.user['_id'], 'story_id': identifier})
        if existing:
            db.favorites.delete_one({'_id': existing['_id'], 'user_id': g.user['_id']})
            flash('已取消收藏。', 'success')
        else:
            if not db.stories.find_one({'_id': identifier, 'status': 'published'}):
                abort(404)
            try:
                db.favorites.insert_one({'user_id': g.user['_id'], 'story_id': identifier,
                                         'created_at': now()})
            except DuplicateKeyError:
                pass
            flash('已加入收藏。', 'success')
        return redirect(url_for('detail', story_id=story_id))

    @app.get('/favorites')
    @login_required
    def favorites():
        identifiers = [x['story_id'] for x in db.favorites.find({'user_id': g.user['_id']})]
        stories = list(db.stories.find({'_id': {'$in': identifiers}, 'status': 'published'}).sort('_id', 1))
        return render_template('favorites.html', stories=stories)

    @app.get('/manage')
    @admin_required
    def manage():
        state = request.args.get('status', '')
        query = {'status': state} if state in ['draft', 'published', 'deleted'] else {}
        return render_template('manage.html', stories=list(db.stories.find(query).sort('updated_at', -1)),
                               counts={s: db.stories.count_documents({'status': s})
                                       for s in ['draft', 'published', 'deleted']},
                               pending=db.ai_results.count_documents({'review_status': 'pending'}))

    @app.get('/manage/<story_id>')
    @admin_required
    def manage_detail(story_id):
        story = db.stories.find_one({'_id': oid(story_id)})
        if not story:
            abort(404)
        results = list(db.ai_results.find({'story_id': story['_id']}).sort('generated_at', -1))
        return render_template('review.html', story=story, results=results)

    @app.post('/manage/<story_id>/status')
    @admin_required
    def set_status(story_id):
        story = db.stories.find_one({'_id': oid(story_id)})
        if not story:
            abort(404)
        action = request.form.get('action')
        if action == 'publish':
            if story['status'] == 'deleted':
                flash('请先恢复草稿，再发布。', 'error')
            elif not story.get('summary', '').strip() or story['story_type'] not in TYPES:
                flash('发布前须补充摘要，并选择正式分类。', 'error')
            elif any(not (Path(app.config['UPLOAD_FOLDER']) / x['path']).is_file()
                     for x in story.get('images', [])):
                flash('图片文件不存在，不能发布。', 'error')
            else:
                db.stories.update_one({'_id': story['_id']}, {'$set': {'status': 'published', 'updated_at': now()}})
                flash('故事已发布。', 'success')
        elif action in ['delete', 'restore', 'draft']:
            target = 'deleted' if action == 'delete' else 'draft'
            db.stories.update_one({'_id': story['_id']}, {'$set': {'status': target, 'updated_at': now()}})
            flash('状态已更新。', 'success')
        else:
            abort(400)
        return redirect(url_for('manage'))

    @app.post('/manage/<story_id>/process')
    @admin_required
    def generate(story_id):
        story = db.stories.find_one({'_id': oid(story_id), 'status': {'$ne': 'deleted'}})
        if not story:
            abort(404)
        try:
            output, model = process(story)
            db.ai_results.insert_one({'story_id': story['_id'], 'model_name': model,
                'prompt_version': 'v1', 'input_hash': input_hash(story), 'output': output,
                'review_status': 'pending', 'generated_at': now()})
            flash('处理结果已生成，请对照原文审核。', 'success')
        except Exception:
            # Do not expose keys or upstream response bodies in pages or logs.
            flash('处理失败，请检查模型配置或输入内容。未保存不合格结果。', 'error')
        return redirect(url_for('manage_detail', story_id=story_id))

    @app.post('/manage/results/<result_id>/review')
    @admin_required
    def review_result(result_id):
        result = db.ai_results.find_one({'_id': oid(result_id)})
        if not result:
            abort(404)
        try:
            if request.form.get('action') == 'reject':
                changed = db.ai_results.update_one({'_id': result['_id'], 'review_status': 'pending'},
                    {'$set': {'review_status': 'rejected', 'reviewed_at': now()}})
                if not changed.modified_count:
                    raise ValueError('该结果已审核，不能重复操作。')
            elif request.form.get('action') == 'accept':
                output = validate_output({'story_type': request.form.get('story_type'),
                    'summary': request.form.get('summary'),
                    'theme_keywords': split_values(request.form.get('theme_keywords'))})

                def adopt(s):
                    fresh = db.ai_results.find_one({'_id': result['_id'], 'review_status': 'pending'}, session=s)
                    story = db.stories.find_one({'_id': result['story_id'], 'status': {'$ne': 'deleted'}}, session=s)
                    if not fresh or not story:
                        raise ValueError('结果已审核或故事已删除。')
                    if input_hash(story) != fresh['input_hash']:
                        raise ValueError('输入内容已变化，请重新生成结果。')
                    db.stories.update_one({'_id': story['_id']}, {'$set': {**output, 'updated_at': now()}}, session=s)
                    db.ai_results.update_one({'_id': fresh['_id']}, {'$set': {'output': output,
                        'review_status': 'accepted', 'reviewed_at': now()}}, session=s)
                with client.start_session() as mongo_session:
                    mongo_session.with_transaction(adopt)
            else:
                abort(400)
            flash('审核已完成。', 'success')
        except ValueError as error:
            flash(str(error), 'error')
        except PyMongoError:
            flash('数据库事务未完成；请确认MongoDB副本集正常。', 'error')
        return redirect(url_for('manage_detail', story_id=str(result['story_id'])))

    @app.post('/manage/<story_id>/image')
    @admin_required
    def upload_image(story_id):
        story = db.stories.find_one({'_id': oid(story_id), 'status': {'$ne': 'deleted'}})
        if not story:
            abort(404)
        uploaded = request.files.get('image')
        try:
            source = canonical_url(request.form.get('source_url'))
            if not uploaded:
                raise ValueError('请选择图片。')
            with Image.open(uploaded.stream) as im:
                fmt = im.format
                if fmt not in ['PNG', 'JPEG', 'WEBP'] or im.width * im.height > 20_000_000:
                    raise ValueError('只接受PNG、JPEG、WEBP图片，最多2000万像素。')
                im.verify()
            uploaded.stream.seek(0)
            name = uuid.uuid4().hex + {'PNG': '.png', 'JPEG': '.jpg', 'WEBP': '.webp'}[fmt]
            path = Path(app.config['UPLOAD_FOLDER']) / name
            uploaded.save(path)
            image = {'path': name, 'source_url': source,
                     'caption': request.form.get('caption', '').strip()}
            try:
                db.stories.update_one({'_id': story['_id']}, {'$push': {'images': image},
                    '$set': {'updated_at': now(), 'status': 'draft'}})
            except Exception:
                path.unlink(missing_ok=True)
                raise
            flash('图片已保存，故事已转为待核对的草稿。', 'success')
        except (ValueError, OSError) as error:
            flash(str(error), 'error')
        return redirect(url_for('manage_detail', story_id=story_id))

    @app.get('/images/<name>')
    def image(name):
        if not re.fullmatch(r'[a-f0-9]{32}\.(png|jpg|webp)', name):
            abort(404)
        query = {'images.path': name}
        if not g.user or g.user['role'] != 'admin':
            query['status'] = 'published'
        if not db.stories.find_one(query):
            abort(404)
        return send_from_directory(app.config['UPLOAD_FOLDER'], name)

    @app.route('/manage/import', methods=['GET', 'POST'])
    @admin_required
    def import_data():
        report = None
        if request.method == 'POST':
            file = request.files.get('file')
            report = {'added': 0, 'duplicate': 0, 'errors': []}
            try:
                if not file:
                    raise ValueError('请选择CSV或Excel文件。')
                if file.filename.lower().endswith('.csv'):
                    rows = list(csv.DictReader(io.StringIO(file.read().decode('utf-8-sig'))))
                elif file.filename.lower().endswith('.xlsx'):
                    book = load_workbook(file, read_only=True, data_only=True)
                    raw = book.active.iter_rows(values_only=True)
                    headers = [str(v or '').strip() for v in next(raw)]
                    rows = [dict(zip(headers, row)) for row in raw if any(v is not None for v in row)]
                    book.close()
                else:
                    raise ValueError('仅支持UTF-8 CSV和.xlsx文件。')
                if len(rows) > 2000:
                    raise ValueError('每次最多导入2000行。')
                for rowno, row in enumerate(rows, 2):
                    try:
                        row = {str(k).strip(): str(v).strip() if v is not None else ''
                               for k, v in row.items() if k is not None}
                        row['content_kind'] = row.get('content_kind') or '故事正文'
                        row['story_type'] = row.get('story_type') or '待分类'
                        doc = story_from_form(row)
                        doc['status'] = 'draft'
                        db.stories.insert_one(doc)
                        report['added'] += 1
                    except DuplicateKeyError:
                        report['duplicate'] += 1
                    except (ValueError, PyMongoError) as error:
                        message = str(error) if isinstance(error, ValueError) else '数据库拒绝此记录。'
                        report['errors'].append({'row': rowno, 'reason': message})
            except (ValueError, UnicodeDecodeError, StopIteration, KeyError, zipfile.BadZipFile):
                report['errors'].append({'row': 0, 'reason': '无法读取文件，请检查格式、编码和表头。'})
        return render_template('import.html', report=report)

    class ProtectedIndex(AdminIndexView):
        def is_accessible(self):
            return bool(g.user and g.user['role'] == 'admin')

        def inaccessible_callback(self, name, **kwargs):
            return redirect(url_for('login')) if not g.user else abort(403)

        @expose('/')
        def index(self):
            return redirect(url_for('manage'))

    class StoryView(ModelView):
        form = StoryForm
        column_list = ['title', 'story_type', 'status', 'updated_at']
        column_labels = {'title': '故事名称', 'story_type': '类型', 'status': '状态', 'updated_at': '修改时间'}
        column_searchable_list = ['title']
        column_sortable_list = ['title', 'updated_at']
        action_disallowed_list = ['delete']
        page_size = 15

        def is_accessible(self):
            return bool(g.user and g.user['role'] == 'admin')

        def inaccessible_callback(self, name, **kwargs):
            return redirect(url_for('login')) if not g.user else abort(403)

        def edit_form(self, obj):
            data = dict(obj)
            for key in ['regions', 'ethnic_groups', 'main_characters', 'theme_keywords']:
                data[key] = ', '.join(obj.get(key, []))
            data['site_name'] = obj['source']['site_name']
            data['source_url'] = obj['source']['url']
            return self.form(request.form if request.method == 'POST' else None, data=data)

        def create_model(self, form):
            try:
                doc = story_from_form(form.data)
                doc['_id'] = self.coll.insert_one(doc).inserted_id
                return doc
            except (ValueError, DuplicateKeyError) as error:
                flash('该来源同名记录已存在。' if isinstance(error, DuplicateKeyError) else str(error), 'error')
                return False

        def update_model(self, form, model):
            try:
                doc = story_from_form(form.data, model)
                self.coll.update_one({'_id': model['_id']}, {'$set': doc})
                return True
            except (ValueError, DuplicateKeyError) as error:
                flash('该来源同名记录已存在。' if isinstance(error, DuplicateKeyError) else str(error), 'error')
                return False

        def delete_model(self, model):
            self.coll.update_one({'_id': model['_id']}, {'$set': {'status': 'deleted', 'updated_at': now()}})
            return True

    admin = Admin(app, name='民间故事资源库', index_view=ProtectedIndex())
    admin.add_view(StoryView(db.stories, name='故事管理', endpoint='stories'))
    return app


if __name__ == '__main__':
    app = create_app()
    app.run(host='0.0.0.0', port=int(os.environ.get('PORT', 5000)), debug=False)
